#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MarkGate 人材紹介エージェント様向けご提案資料 — PowerPoint(.pptx) 生成スクリプト

sales/index.html（A4横・15枚）と同じ構成・配色（Deep Navy × Gold）の
編集可能なネイティブPPTXを生成します。

使い方:
    pip install python-pptx
    python3 sales/generate_pptx.py [出力パス.pptx]

フォントは PowerPoint 標準の「游明朝」（見出し）/「游ゴシック」（本文）、
欧文は Garamond を指定しています（未インストール環境では自動代替）。
"""
import sys

from pptx import Presentation
from pptx.util import Mm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- 配色（style.css のトークンに対応） ----------
NAVY = RGBColor.from_string("0A0E1A")
NAVY7 = RGBColor.from_string("141B33")
NAVY_BOX = RGBColor.from_string("131A30")   # ダーク面のカード
INK = RGBColor.from_string("1A1D29")
INK_SOFT = RGBColor.from_string("3A3F4D")
GOLD = RGBColor.from_string("C6A15B")
GOLD_L = RGBColor.from_string("E2C88A")
GOLD_T = RGBColor.from_string("7A6030")     # 小サイズ文字用の濃いゴールド
CREAM = RGBColor.from_string("F8F6F1")
CREAM2 = RGBColor.from_string("F1EDE3")
PAPER = RGBColor.from_string("FFFFFF")
HL_BG = RGBColor.from_string("FDF9F0")      # 強調カードの背景
HL_BG2 = RGBColor.from_string("FBF4E6")     # フロー強調
MUTED = RGBColor.from_string("5F6672")
LINE_GOLD = RGBColor.from_string("DECBA2")  # 罫線（ゴールド系）
LINE_SOFT = RGBColor.from_string("D9D9DE")  # 罫線（グレー系）
WHITE = RGBColor.from_string("FFFFFF")
DARK_BODY = RGBColor.from_string("B9BFCC")  # ダーク面の本文
DARK_SUB = RGBColor.from_string("8E96A8")   # ダーク面の注記

SERIF_LAT, SERIF_EA = "Garamond", "游明朝"
SANS_LAT, SANS_EA = "Yu Gothic", "游ゴシック"

FOOTER_TEXT = "人材紹介エージェント様向けご提案資料"


def _set_fonts(run, serif=False, latin=None, ea=None):
    lat = latin or (SERIF_LAT if serif else SANS_LAT)
    east = ea or (SERIF_EA if serif else SANS_EA)
    run.font.name = lat
    rPr = run._r.get_or_add_rPr()
    ea_el = rPr.find(qn("a:ea"))
    if ea_el is None:
        ea_el = rPr.makeelement(qn("a:ea"), {})
        latin_el = rPr.find(qn("a:latin"))
        if latin_el is not None:
            latin_el.addnext(ea_el)
        else:
            rPr.append(ea_el)
    ea_el.set("typeface", east)


def _style_run(run, text, size=9, color=INK, bold=False, serif=False, spc=None):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    _set_fonts(run, serif=serif)
    if spc:
        run._r.get_or_add_rPr().set("spc", str(int(spc * 100)))


def text_box(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    """paras: [(runs, opts)] / runs: [(text, style_dict)] / opts: align, spacing, after"""
    tb = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, (runs, opts) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", PP_ALIGN.LEFT)
        p.line_spacing = opts.get("spacing", 1.25)
        if opts.get("after"):
            p.space_after = Pt(opts["after"])
        if opts.get("before"):
            p.space_before = Pt(opts["before"])
        for text, st in runs:
            _style_run(p.add_run(), text, **st)
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, dash=None,
         shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Mm(x), Mm(y), Mm(w), Mm(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
        if dash:
            ln = sp.line._get_or_add_ln()
            d = ln.makeelement(qn("a:prstDash"), {"val": dash})
            ln.append(d)
    sp.shadow.inherit = False
    sp.text_frame.word_wrap = True
    return sp


def add_slide(prs, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 297, 210, fill=(NAVY if dark else CREAM))
    return slide


def footer(slide, page, dark=False):
    c_brand = GOLD_L if dark else GOLD_T
    c_txt = DARK_SUB if dark else MUTED
    text_box(slide, 19, 199.5, 200, 6, [
        ([("MarkGate", dict(size=9, color=c_brand, bold=True, serif=True)),
          ("　" + FOOTER_TEXT, dict(size=7, color=c_txt))], {}),
    ])
    text_box(slide, 250, 199.5, 28, 6, [
        ([(f"{page:02d}", dict(size=9, color=c_txt, serif=True, spc=1))],
         dict(align=PP_ALIGN.RIGHT)),
    ])


def header(slide, num, label, title_runs, dark=False, lead=None):
    """セクション見出し。コンテンツ開始Y(mm)を返す。"""
    text_box(slide, 19, 12, 20, 10, [
        ([(num, dict(size=17, color=GOLD, serif=True))], {}),
    ])
    text_box(slide, 33, 15.2, 200, 6, [
        ([(label, dict(size=7.5, color=(GOLD_L if dark else GOLD_T), spc=3))], {}),
    ])
    n_lines = sum(1 for _ in title_runs)
    text_box(slide, 19, 22, 259, 12 * n_lines, [
        (runs, dict(spacing=1.2)) for runs in title_runs
    ])
    y = 22 + 12 * n_lines
    if lead:
        text_box(slide, 19, y + 1, 250, 10, [
            (lead, dict(spacing=1.4)),
        ])
        y += 10
    rect(slide, 19, y + 2, 14, 0.7, fill=GOLD)
    return y + 8


def t(text, **kw):
    return (text, kw)


def title_r(text, color=INK, gold=False):
    return t(text, size=22, color=(GOLD_T if gold else color), serif=True, bold=True)


def title_r_dark(text, gold=False):
    return t(text, size=22, color=(GOLD_L if gold else WHITE), serif=True, bold=True)


# ============================================================
def build(out_path):
    prs = Presentation()
    prs.slide_width = Mm(297)
    prs.slide_height = Mm(210)

    # ---------- P.01 表紙 ----------
    s = add_slide(prs, dark=True)
    # ゲートモチーフ（ロゴ簡略版）
    try:
        fb = s.shapes.build_freeform(Mm(24), Mm(52), scale=1)
        fb.add_line_segments(
            [(Mm(24), Mm(38)), (Mm(31), Mm(33.5)), (Mm(38), Mm(38)),
             (Mm(38), Mm(52)), (Mm(24), Mm(52))], close=True)
        gate = fb.convert_to_shape()
        gate.fill.background()
        gate.line.color.rgb = GOLD_L
        gate.line.width = Pt(1.4)
        gate.shadow.inherit = False
    except Exception:
        pass
    text_box(s, 24, 60, 240, 7, [
        ([t("人材紹介エージェント様向け ご提案資料", size=10.5, color=GOLD_L, spc=4)], {}),
    ])
    text_box(s, 24, 72, 250, 42, [
        ([t("貴社のエースが、", size=33, color=WHITE, serif=True, bold=True)],
         dict(spacing=1.25)),
        ([t("“指名される”", size=33, color=GOLD_L, serif=True, bold=True),
          t("時代へ。", size=33, color=WHITE, serif=True, bold=True)],
         dict(spacing=1.25)),
    ])
    text_box(s, 24, 118, 250, 18, [
        ([t("トップキャリアアドバイザーだけが登録できる、", size=11.5,
            color=RGBColor.from_string("D5D9E2"))], dict(spacing=1.5)),
        ([t("ハイクラス特化の指名型転職プラットフォーム「MarkGate」のご案内",
            size=11.5, color=RGBColor.from_string("D5D9E2"))], dict(spacing=1.5)),
    ])
    text_box(s, 24, 150, 250, 9, [
        ([t("MarkGate", size=14, color=WHITE, bold=True, serif=True),
          t("　／　", size=9, color=GOLD),
          t("MarkGate株式会社", size=9, color=DARK_BODY, spc=1),
          t("　／　", size=9, color=GOLD),
          t("2026年7月", size=9, color=DARK_BODY, spc=1)], {}),
    ])

    # ---------- P.02 サマリー ----------
    s = add_slide(prs)
    y = header(s, "01", "EXECUTIVE SUMMARY",
               [[title_r("本日お伝えしたい、3つのこと。")]])
    cards = [
        ("Ⅰ", "MarkGateは\nトップアドバイザー限定の\n指名型プラットフォーム",
         "各エージェントの実績上位アドバイザーのみが審査制で登録できる、ハイクラス・エグゼクティブ特化の転職プラットフォームです。"),
        ("Ⅱ", "ハイクラス求職者からの\n「指名」が、貴社に\n直接届く仕組み",
         "求職者がアドバイザーの実績・専門領域を比較して指名する“指名型（プル型）”モデル。温度感の高い接点が、スカウト送信なしで生まれる設計です。"),
        ("Ⅲ", "審査通過そのものが、\n貴社とエースの\nブランドに",
         "「Mark＝一流の証」を掲げる審査制の場に立つこと自体が、アドバイザー個人と貴社の信頼の可視化につながります。"),
    ]
    cw, ch, gap = 82.3, 88, 6
    cy = y + 4
    for i, (num, head, body) in enumerate(cards):
        cx = 19 + i * (cw + gap)
        rect(s, cx, cy, cw, ch, fill=PAPER, line=LINE_SOFT, line_w=0.5)
        rect(s, cx, cy, cw, 0.9, fill=GOLD)
        paras = [([t(num, size=15, color=GOLD, serif=True)], dict(after=6))]
        for ln in head.split("\n"):
            paras.append(([t(ln, size=12.5, color=INK, serif=True, bold=True)],
                          dict(spacing=1.3)))
        paras.append(([t(body, size=8.5, color=INK_SOFT)],
                      dict(spacing=1.45, before=7)))
        text_box(s, cx + 7, cy + 7, cw - 14, ch - 14, paras)
    by = cy + ch + 7
    rect(s, 19, by, 259, 17, fill=NAVY)
    rect(s, 27, by + 5.2, 24, 6.6, line=GOLD_L, line_w=0.75)
    text_box(s, 27, by + 6.6, 24, 4, [
        ([t("本日のご提案", size=7, color=GOLD_L, spc=2)], dict(align=PP_ALIGN.CENTER)),
    ])
    text_box(s, 57, by, 214, 17, [
        ([t("貴社の", size=12, color=WHITE, serif=True),
          t("トップキャリアアドバイザー", size=12, color=GOLD_L, serif=True),
          t("の、MarkGateへのご登録（審査制）をご検討ください。",
            size=12, color=WHITE, serif=True)], {}),
    ], anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 2)

    # ---------- P.03 課題（ダーク） ----------
    s = add_slide(prs, dark=True)
    y = header(s, "02", "BACKGROUND",
               [[title_r_dark("ハイクラス紹介の現場で、")],
                [title_r_dark("いま起きていること。")]], dark=True)
    issues = [
        ("Ⅰ", "スカウト依存の消耗戦",
         "各社が同じデータベースへ大量のスカウトを送り合い、候補者の受信箱は飽和。送信量を増やしても、返信は容易に増えない構造になりつつあります。"),
        ("Ⅱ", "ハイクラス層ほど“売り込み”に反応しない",
         "経営幹部・高度専門職ほど、届く連絡の量ではなく「誰からの連絡か」で動きます。選ぶ立場の求職者に対して、“選ばれるための接点”が不足しています。"),
        ("Ⅲ", "エースの実績が、社外から見えない",
         "どれほど成約実績を積み上げても、トップアドバイザー個人の信頼は社外に可視化されにくく、個人の実力が新しい出会いにつながりません。"),
        ("Ⅳ", "集客コストの上昇",
         "媒体掲載費・スカウト課金・広告費は上昇傾向にあり、決定単価の高いハイクラス領域でも、集客コストが利益を圧迫し始めています。"),
    ]
    bw, bh, gap = 126.5, 44, 6
    for i, (num, head, body) in enumerate(issues):
        bx = 19 + (i % 2) * (bw + gap)
        byy = y + 4 + (i // 2) * (bh + gap)
        rect(s, bx, byy, bw, bh, fill=NAVY_BOX, line=RGBColor.from_string("4A4531"),
             line_w=0.5)
        text_box(s, bx + 7, byy + 6, 10, 10, [
            ([t(num, size=14, color=GOLD_L, serif=True)], {}),
        ])
        text_box(s, bx + 19, byy + 6, bw - 26, bh - 12, [
            ([t(head, size=11.5, color=WHITE, serif=True, bold=True)],
             dict(after=3)),
            ([t(body, size=8.5, color=DARK_BODY)], dict(spacing=1.4)),
        ])
    text_box(s, 19, y + 4 + 2 * (bh + gap) + 2, 259, 6, [
        ([t("※ 本ページは業界環境に関する当社の課題認識を整理したものです。",
            size=7.5, color=DARK_SUB)], {}),
    ])
    footer(s, 3, dark=True)

    # ---------- P.04 発想の転換 ----------
    s = add_slide(prs)
    y = header(s, "03", "PARADIGM SHIFT",
               [[title_r("「探して、送る」から、"), title_r("「選ばれて、届く」", gold=True),
                 title_r("へ。")]])
    colw, colh = 119, 92
    cy = y + 6

    def shift_col(x, new):
        if new:
            rect(s, x, cy, colw, colh, fill=HL_BG, line=GOLD, line_w=1)
        else:
            rect(s, x, cy, colw, colh, fill=PAPER, line=LINE_SOFT, line_w=0.5)

    shift_col(19, False)
    shift_col(19 + colw + 21, True)
    arrow = text_box(s, 19 + colw + 2, cy, 17, colh, [
        ([t("→", size=20, color=GOLD, serif=True)], dict(align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)

    def diagram(x, yy, nodes):
        nx = x
        for i, (label, navy) in enumerate(nodes):
            wdt = 8 + len(label) * 4.2
            rect(s, nx, yy, wdt, 8.5, fill=(NAVY if navy else CREAM),
                 line=(NAVY if navy else LINE_SOFT), line_w=0.5)
            text_box(s, nx, yy + 1.2, wdt, 6, [
                ([t(label, size=8.5, color=(GOLD_L if navy else INK),
                    serif=navy)], dict(align=PP_ALIGN.CENTER)),
            ])
            nx += wdt + 2
            if i < len(nodes) - 1:
                text_box(s, nx - 1, yy + 1.4, 6, 6, [
                    ([t("→", size=9, color=GOLD_T)], {}),
                ])
                nx += 5

    # 左列（従来）
    x = 19
    rect(s, x + 7, cy + 6, 30, 6.4, fill=CREAM2)
    text_box(s, x + 7, cy + 7.5, 30, 4, [
        ([t("従来のスカウト型", size=7, color=MUTED, spc=1.5)],
         dict(align=PP_ALIGN.CENTER)),
    ])
    text_box(s, x + 7, cy + 15, colw - 14, 8, [
        ([t("エージェントが、求職者を探す", size=13, color=INK, serif=True, bold=True)], {}),
    ])
    diagram(x + 7, cy + 25, [("アドバイザー", False), ("スカウト送信", False), ("求職者", False)])
    text_box(s, x + 7, cy + 37, colw - 14, 14, [
        ([t("送信の「量」が接点を左右するプッシュ型。返信率に左右され、面談開始時点の温度感も読みにくい構造です。",
            size=8.5, color=INK_SOFT)], dict(spacing=1.4)),
    ])
    bl = ["大量送信・文面作成の工数が常時発生",
          "候補者は比較検討前で、温度感が低いことも",
          "成果が送信量と媒体費に依存"]
    text_box(s, x + 7, cy + 55, colw - 14, 30, [
        ([t("－ ", size=8.5, color=MUTED), t(b, size=8.5, color=INK_SOFT)],
         dict(spacing=1.5)) for b in bl
    ])

    # 右列（MarkGate）
    x = 19 + colw + 21
    rect(s, x + 7, cy + 6, 34, 6.4, fill=NAVY)
    text_box(s, x + 7, cy + 7.5, 34, 4, [
        ([t("MarkGateの指名型", size=7, color=GOLD_L, spc=1.5)],
         dict(align=PP_ALIGN.CENTER)),
    ])
    text_box(s, x + 7, cy + 15, colw - 14, 8, [
        ([t("求職者が、アドバイザーを選ぶ", size=13, color=INK, serif=True, bold=True)], {}),
    ])
    diagram(x + 7, cy + 25, [("求職者", False), ("MarkGate", True), ("指名が届く", False)])
    text_box(s, x + 7, cy + 37, colw - 14, 14, [
        ([t("求職者が実績・専門領域を比較したうえで指名するプル型。「あなたに任せたい」という状態から支援が始まります。",
            size=8.5, color=INK_SOFT)], dict(spacing=1.4)),
    ])
    br = ["指名＝比較検討を終えた、温度感の高い接点",
          "プロフィールが、貴社エースの常設の窓口として機能する設計",
          "実力と専門性が、そのまま機会に変わる"]
    text_box(s, x + 7, cy + 55, colw - 14, 32, [
        ([t("－ ", size=8.5, color=GOLD_T), t(b, size=8.5, color=INK_SOFT)],
         dict(spacing=1.5)) for b in br
    ])
    footer(s, 4)

    # ---------- P.05 MarkGateとは ----------
    s = add_slide(prs)
    y = header(s, "04", "ABOUT MARKGATE", [[title_r("MarkGateとは。")]],
               lead=[t("各人材紹介エージェントの実績トップクラスのキャリアアドバイザーだけが集う、ハイクラス層のための転職プラットフォームです。",
                       size=10, color=INK_SOFT)])
    pillars = [
        ("01", "トップアドバイザー限定\n（審査制）",
         "各エージェントの実績上位者のみが、経歴・成約実績・専門性を基準とした独自審査を通過して登録。一定水準のアドバイザーだけが集う環境を担保します。"),
        ("02", "ハイクラス・\nエグゼクティブ特化",
         "経営幹部・管理職・高度専門職、年収800万円以上のハイクラス層に特化。決定単価の高い領域に絞った、質の高いマッチングの場です。"),
        ("03", "アドバイザーを\n“選べる・指名できる”",
         "求職者はアドバイザーの専門領域・実績・支援スタンスを比較し、自分に最適な伴走者を指名。指名はアドバイザーへ直接届きます。"),
    ]
    cw, ch, gap = 82.3, 74, 6
    cy = y + 3
    for i, (num, head, body) in enumerate(pillars):
        cx = 19 + i * (cw + gap)
        rect(s, cx, cy, cw, ch, fill=PAPER, line=LINE_SOFT, line_w=0.5)
        rect(s, cx + 7, cy + ch - 1.2, cw - 14, 0.8, fill=GOLD_L)
        paras = [([t(num, size=15, color=GOLD, serif=True)], dict(after=4))]
        for ln in head.split("\n"):
            paras.append(([t(ln, size=12, color=INK, serif=True, bold=True)],
                          dict(spacing=1.3)))
        paras.append(([t(body, size=8.5, color=INK_SOFT)],
                      dict(spacing=1.45, before=5)))
        text_box(s, cx + 7, cy + 6, cw - 14, ch - 12, paras)
    by = cy + ch + 7
    rect(s, 19, by, 259, 22, fill=CREAM2)
    rect(s, 19, by, 1.2, 22, fill=GOLD)
    text_box(s, 27, by, 245, 22, [
        ([t("“", size=11.5, color=INK, serif=True),
          t("Mark", size=11.5, color=GOLD_T, serif=True, bold=True),
          t("（一流の証）”を持つ者だけが立つ、キャリアの“", size=11.5, color=INK, serif=True),
          t("Gate", size=11.5, color=GOLD_T, serif=True, bold=True),
          t("（門）”。", size=11.5, color=INK, serif=True)], dict(spacing=1.5)),
        ([t("—— 私たちは、実力あるアドバイザーが正当に選ばれる市場をつくります。",
            size=11.5, color=INK, serif=True)], dict(spacing=1.5)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 5)

    # ---------- P.06 仕組み ----------
    s = add_slide(prs)
    y = header(s, "05", "HOW IT WORKS", [[title_r("サービスの仕組み。")]])

    def lane(yy, tag, navy_tag, steps, hl_idx):
        if navy_tag:
            rect(s, 19, yy, 30, 6.4, fill=NAVY)
            text_box(s, 19, yy + 1.5, 30, 4, [
                ([t(tag, size=7, color=GOLD_L, spc=1.5)], dict(align=PP_ALIGN.CENTER)),
            ])
        else:
            rect(s, 19, yy, 40, 6.4, fill=CREAM2, line=LINE_GOLD, line_w=0.5)
            text_box(s, 19, yy + 1.5, 40, 4, [
                ([t(tag, size=7, color=GOLD_T, spc=1.5)], dict(align=PP_ALIGN.CENTER)),
            ])
        sw, sh = 49.4, 30
        for i, (head, body) in enumerate(steps):
            sx = 19 + i * (sw + 3)
            hl = (i == hl_idx)
            rect(s, sx, yy + 9, sw, sh, fill=(HL_BG2 if hl else PAPER),
                 line=(GOLD if hl else LINE_SOFT), line_w=0.75 if hl else 0.5)
            text_box(s, sx + 4, yy + 12, sw - 8, sh - 6, [
                ([t(f"STEP 0{i+1}", size=7.5, color=GOLD_T, serif=True, spc=1)],
                 dict(after=2)),
                ([t(head, size=9.5, color=INK, serif=True, bold=True)],
                 dict(after=1.5, spacing=1.2)),
                ([t(body, size=7.2, color=INK_SOFT)], dict(spacing=1.35)),
            ])
            if i < len(steps) - 1:
                text_box(s, sx + sw - 1.2, yy + 9 + sh / 2 - 3, 6, 6, [
                    ([t("›", size=12, color=GOLD)], {}),
                ])

    lane(y + 2, "求職者の動き", True, [
        ("会員登録（審査制）", "ハイクラス特化のため、経歴をもとに審査を実施。"),
        ("アドバイザーを検索・比較", "業界・職種・対応年収・支援スタンスで絞り込み。"),
        ("“この人”を指名", "実績・専門領域に納得したうえで、指名メッセージを送信。"),
        ("面談・求人提案", "市場価値の見極めと、上位ポジションの提案。"),
        ("選考・内定・入社", "書類・面接対策から条件交渉まで一気通貫。"),
    ], 2)
    lane(y + 50, "貴社アドバイザーの動き", False, [
        ("登録申請", "実績上位のアドバイザーを対象にご申請。"),
        ("審査", "成約実績・専門性・経歴を基準に審査。"),
        ("プロフィール掲載", "専門領域・実績・支援スタンスを掲載。"),
        ("指名を受領", "求職者からの指名がメッセージ付きで届く。"),
        ("支援開始", "面談を設定し、通常どおりの紹介業務へ。"),
    ], 3)
    text_box(s, 19, y + 98, 259, 10, [
        ([t("※ 指名後の転職支援・企業への紹介は、従来どおり貴社（登録アドバイザー）の紹介業務として実施いただきます。MarkGateは求職者との質の高い接点を提供します。",
            size=7.5, color=MUTED)], dict(spacing=1.4)),
    ])
    footer(s, 6)

    # ---------- P.07 メリット ----------
    s = add_slide(prs)
    y = header(s, "06", "BENEFITS FOR AGENCIES",
               [[title_r("貴社にとっての、4つのメリット。")]])
    merits = [
        ("1", "温度感の高い「指名」から面談が始まる",
         "求職者は複数のアドバイザーを比較・納得したうえで指名します。「まず話を聞いてみたい」ではなく「あなたに任せたい」から始まる面談により、支援効率の向上が期待できます。"),
        ("2", "スカウト課金に依存しない、新しい接点",
         "登録・掲載は審査制。掲載そのものが求職者への露出となり、大量送信の工数や媒体課金に依存しない、継続的なハイクラス接点の構築を目指せます。"),
        ("3", "「審査通過」が、エースと貴社の証明になる",
         "トップアドバイザー限定の場に立つこと自体がブランドです。エース個人の市場価値の可視化は、貴社の対外的な信頼・採用力・社内のモチベーションにも波及します。"),
        ("4", "決定単価の高い案件に集中できる",
         "会員は年収800万円以上を想定したハイクラス・エグゼクティブ層。トップアドバイザーの時間を、最も価値の高い支援に集中させることができます。"),
    ]
    bw, bh, gap = 126.5, 46, 6
    for i, (num, head, body) in enumerate(merits):
        bx = 19 + (i % 2) * (bw + gap)
        byy = y + 4 + (i // 2) * (bh + gap)
        rect(s, bx, byy, bw, bh, fill=PAPER, line=LINE_SOFT, line_w=0.5)
        rect(s, bx + 6, byy + 6, 10, 10, line=GOLD, line_w=0.75,
             shape=MSO_SHAPE.OVAL)
        text_box(s, bx + 6, byy + 8, 10, 6, [
            ([t(num, size=11, color=GOLD_T, serif=True)], dict(align=PP_ALIGN.CENTER)),
        ])
        text_box(s, bx + 20, byy + 6, bw - 27, bh - 11, [
            ([t(head, size=11.5, color=INK, serif=True, bold=True)], dict(after=2.5)),
            ([t(body, size=8.5, color=INK_SOFT)], dict(spacing=1.4)),
        ])
    text_box(s, 19, y + 4 + 2 * (bh + gap) + 2, 259, 6, [
        ([t("※ 効果に関する記述はサービス設計に基づく想定であり、成果を保証するものではありません。",
            size=7.5, color=MUTED)], {}),
    ])
    footer(s, 7)

    # ---------- P.08 チャネル比較 ----------
    s = add_slide(prs)
    y = header(s, "07", "CHANNEL COMPARISON",
               [[title_r("既存チャネルと競合せず、"), title_r("補完", gold=True),
                 title_r("します。")]])
    cols = [34, 70, 70, 85]
    rows = [
        ("", "スカウト媒体", "求人広告・自社集客", "MarkGate"),
        ("接点の起点", "アドバイザーからの送信（プッシュ型）", "求人・企業情報への応募",
         "求職者からの「指名」（プル型）"),
        ("求職者の温度感", "比較検討前。返信後の見極めが必要", "案件起点。担当者は選べない",
         "比較・納得済み。「任せたい」状態から開始"),
        ("可視化されるもの", "求人案件・スカウト文面", "会社・求人情報",
         "アドバイザー個人の実績・専門性"),
        ("主な工数・費用", "送信工数＋データベース利用・課金", "掲載費・広告運用",
         "プロフィール整備＋月額費用（開始キャンペーンにより1年目無料※）"),
        ("対象層", "媒体登録者全般", "広く一般",
         "ハイクラス・エグゼクティブ（年収800万円以上を想定）"),
    ]
    row_h = [10, 13, 13, 13, 15, 13]
    ty = y + 3
    for r, row in enumerate(rows):
        tx = 19
        for c, cell in enumerate(row):
            w = cols[c]
            if r == 0:
                fill = GOLD if c == 3 else NAVY
                color = NAVY if c == 3 else WHITE
                rect(s, tx, ty, w, row_h[0], fill=fill)
                if cell:
                    text_box(s, tx + 3, ty, w - 6, row_h[0],
                             [([t(cell, size=9.5, color=color, serif=True,
                                  bold=True)], {})],
                             anchor=MSO_ANCHOR.MIDDLE)
            else:
                if c == 0:
                    fill, color, size, bold = CREAM2, INK, 8.5, True
                elif c == 3:
                    fill, color, size, bold = HL_BG, INK, 8.5, False
                else:
                    fill, color, size, bold = PAPER, INK_SOFT, 8.5, False
                rect(s, tx, ty, w, row_h[r], fill=fill, line=LINE_SOFT, line_w=0.4)
                text_box(s, tx + 3, ty, w - 6, row_h[r],
                         [([t(cell, size=size, color=color, bold=bold)],
                           dict(spacing=1.25))],
                         anchor=MSO_ANCHOR.MIDDLE)
            tx += w
        ty += row_h[r]
    # MarkGate列の強調枠
    rect(s, 19 + sum(cols[:3]), y + 3, cols[3], sum(row_h), line=GOLD, line_w=1)
    text_box(s, 19, ty + 3, 259, 10, [
        ([t("※ 各チャネルの特徴は一般的な傾向を整理したものです。料金の詳細は P.11 をご覧ください。既存のスカウト・広告チャネルと併用いただけます。",
            size=7.5, color=MUTED)], dict(spacing=1.4)),
    ])
    footer(s, 8)

    # ---------- P.09 審査 ----------
    s = add_slide(prs)
    y = header(s, "08", "SCREENING", [[title_r("登録アドバイザーの審査について。")]],
               lead=[t("「トップアドバイザーだけが集う」というプラットフォームの価値は、審査によって守られます。だからこそ、通過したアドバイザーの掲載には意味が生まれます。",
                       size=10, color=INK_SOFT)])
    crits = [
        ("成約実績", "ハイクラス帯（年収800万円以上を想定）での決定実績・成約の質を重視します。"),
        ("専門性", "特定の業界・職種・レイヤーにおける知見の深さと、非公開・上位ポジションへの接続力。"),
        ("経歴・所属での位置づけ", "キャリアアドバイザーとしての経験と、所属エージェントにおける実績上位であること。"),
        ("支援スタンス", "求職者のキャリアに長期目線で向き合う姿勢。プラットフォームの品質をともに守れる方。"),
    ]
    text_box(s, 19, y + 2, 100, 5, [
        ([t("主な審査基準", size=8, color=GOLD_T, spc=2.5)], {}),
    ])
    cy = y + 9
    for head, body in crits:
        rect(s, 19, cy, 140, 19, fill=PAPER, line=LINE_SOFT, line_w=0.5)
        rect(s, 19, cy, 1, 19, fill=GOLD)
        text_box(s, 25, cy + 2.5, 130, 15, [
            ([t(head, size=10.5, color=INK, serif=True, bold=True)], dict(after=1)),
            ([t(body, size=8, color=INK_SOFT)], dict(spacing=1.3)),
        ])
        cy += 22.5
    steps = [
        ("登録申請", "貴社の対象アドバイザーについて、申請フォームからご提出ください。"),
        ("書類審査", "経歴・実績・専門領域を確認します。"),
        ("面談", "支援スタンス・専門性について、オンラインでお話を伺います。"),
        ("掲載開始", "プロフィールを作成し、公開。指名の受付が始まります。"),
    ]
    sx = 172
    text_box(s, sx, y + 2, 100, 5, [
        ([t("審査の流れ", size=8, color=GOLD_T, spc=2.5)], {}),
    ])
    cy = y + 9
    for i, (head, body) in enumerate(steps):
        rect(s, sx, cy + 0.5, 5, 5, line=GOLD, line_w=0.75, shape=MSO_SHAPE.OVAL)
        if i < len(steps) - 1:
            rect(s, sx + 2.3, cy + 6, 0.35, 14.5, fill=LINE_GOLD)
        text_box(s, sx + 9, cy, 97, 18, [
            ([t(head, size=10.5, color=INK, serif=True, bold=True)], dict(after=1)),
            ([t(body, size=8, color=INK_SOFT)], dict(spacing=1.3)),
        ])
        cy += 20.5
    text_box(s, sx, cy + 1, 106, 8, [
        ([t("※ 審査基準の詳細・提出書類は、お問い合わせ後に個別にご案内します。",
            size=7.5, color=MUTED)], dict(spacing=1.35)),
    ])
    footer(s, 9)

    # ---------- P.10 プロフィール掲載イメージ ----------
    s = add_slide(prs)
    y = header(s, "09", "PROFILE",
               [[title_r("プロフィールが、エースの"),
                 title_r("“指名される窓口”", gold=True), title_r("に。")]])
    points = [
        ("経歴・支援実績", "を構造化して掲載。定量・定性の両面で、アドバイザーの実力を伝えます。"),
        ("専門領域（業界／職種／対応年収帯）", "で検索・絞り込みの対象に。得意分野の求職者と出会えます。"),
        ("支援スタンス", "（伴走型・提案型など）も明示。ミスマッチの少ない指名につながります。"),
        ("指名にはメッセージが添付", "され、求職者の意図・温度感を把握したうえで面談に入れます。"),
        ("プロフィールは掲載後も更新可能", "。実績を積むほど、指名の入口が強くなります。"),
    ]
    text_box(s, 19, y + 6, 118, 90, [
        ([t("－ ", size=9.5, color=GOLD_T),
          t(b, size=9.5, color=INK, bold=True),
          t(rest, size=9.5, color=INK_SOFT)],
         dict(spacing=1.5, after=6)) for b, rest in points
    ])
    # サンプルカード
    cx, cy2, cw2, ch2 = 150, y + 3, 128, 92
    rect(s, cx, cy2, cw2, ch2, fill=PAPER, line=LINE_GOLD, line_w=0.75)
    rect(s, cx + 7, cy2 + 6, 15, 15, fill=NAVY7, shape=MSO_SHAPE.OVAL)
    text_box(s, cx + 7, cy2 + 9.5, 15, 8, [
        ([t("M", size=13, color=GOLD_L, serif=True)], dict(align=PP_ALIGN.CENTER)),
    ])
    text_box(s, cx + 26, cy2 + 5.5, cw2 - 33, 18, [
        ([t("掲載イメージ（サンプル）", size=11.5, color=INK, serif=True, bold=True)],
         dict(after=1)),
        ([t("シニアキャリアアドバイザー ｜ ○○エージェント所属", size=7.5, color=MUTED)],
         dict(after=2)),
        ([t(" 審査通過 ", size=7, color=GOLD_T),
          t("　", size=7),
          t(" 経営幹部・CxO ", size=7, color=GOLD_T),
          t("　", size=7),
          t(" 年収1,000万円〜 ", size=7, color=GOLD_T)], {}),
    ])
    rect(s, cx + 7, cy2 + 25, cw2 - 14, 0.4, fill=LINE_SOFT)
    prof_rows = [
        ("専門領域", "製造業・メーカー ／ 経営企画・事業責任者クラス"),
        ("支援実績", "ハイクラス帯の決定実績、リピート・紹介率などを掲載"),
        ("支援スタンス", "中長期のキャリア戦略から逆算する伴走型"),
        ("経歴", "キャリア・得意分野の背景を記載"),
    ]
    ry = cy2 + 29
    for k, v in prof_rows:
        text_box(s, cx + 7, ry, 26, 6, [
            ([t(k, size=8, color=MUTED)], {}),
        ])
        text_box(s, cx + 35, ry, cw2 - 42, 6, [
            ([t(v, size=8, color=INK)], {}),
        ])
        ry += 8.5
    rect(s, cx + 7, ry + 2, cw2 - 14, 9, fill=GOLD)
    text_box(s, cx + 7, ry + 4.2, cw2 - 14, 5, [
        ([t("このアドバイザーを指名する", size=8.5, color=NAVY, bold=True, spc=1.5)],
         dict(align=PP_ALIGN.CENTER)),
    ])
    text_box(s, 19, y + 100, 259, 6, [
        ([t("※ 上記はイメージです。実際の掲載項目・レイアウトはプラットフォームの仕様に準じます。",
            size=7.5, color=MUTED)], {}),
    ])
    footer(s, 10)

    # ---------- P.11 料金 ----------
    s = add_slide(prs)
    y = header(s, "10", "PRICING", [[title_r("ご利用料金。")]],
               lead=[t("開始キャンペーン期間中にご参画いただくと、1年目の月額費用は無料です。",
                       size=10, color=INK_SOFT)])
    prices = [
        ("初期費用", "0", "円", "登録申請・審査に費用はかかりません。", False),
        ("月額費用（1年目）", "0", "円",
         "通常 月額5万円のところ、開始キャンペーンにより現在は無料でご利用いただけます。", True),
        ("月額費用（2年目以降）", "5", "万円／月",
         "2年目以降は、月額5万円でご利用いただけます。", False),
    ]
    cw, ch, gap = 82.3, 62, 6
    cy = y + 4
    for i, (label, num, unit, body, hl) in enumerate(prices):
        cx = 19 + i * (cw + gap)
        rect(s, cx, cy, cw, ch, fill=(HL_BG if hl else PAPER),
             line=(GOLD if hl else LINE_SOFT), line_w=1 if hl else 0.5)
        text_box(s, cx + 6, cy + 8, cw - 12, ch - 14, [
            ([t(label, size=8.5, color=GOLD_T, spc=1.5)],
             dict(align=PP_ALIGN.CENTER, after=6)),
            ([t(num, size=24, color=INK, serif=True, bold=True),
              t(" " + unit, size=11, color=INK_SOFT, serif=True)],
             dict(align=PP_ALIGN.CENTER, after=5)),
            ([t(body, size=8.5, color=INK_SOFT)],
             dict(align=PP_ALIGN.CENTER, spacing=1.45)),
        ])
    by = cy + ch + 8
    rect(s, 19, by, 259, 20, fill=CREAM2, line=GOLD, line_w=0.75, dash="dash")
    text_box(s, 27, by, 243, 20, [
        ([t("【開始キャンペーンについて】", size=9, color=GOLD_T, bold=True),
          t("適用条件・期間の詳細は、お申し込み時にご案内します。キャンペーンは予告なく終了する場合があります。正式なご契約条件は、契約書面にてご案内します。",
            size=9, color=INK_SOFT)], dict(spacing=1.5)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 11)

    # ---------- P.12 導入までの流れ ----------
    s = add_slide(prs)
    y = header(s, "11", "ONBOARDING", [[title_r("導入までの流れ。")]])
    onb = [
        ("お問い合わせ", "本資料の内容について、オンラインで詳しくご説明します。"),
        ("登録候補の選定", "貴社の実績上位アドバイザーから、登録候補をご選定ください。"),
        ("登録申請・審査", "書類審査と面談を実施します（無料・料金はP.11）。"),
        ("プロフィール掲載", "掲載内容を一緒に磨き込み、公開します。"),
        ("指名受領・支援開始", "求職者からの指名を受け、面談・支援を開始します。"),
    ]
    sw, sh = 49.4, 44
    cy = y + 4
    for i, (head, body) in enumerate(onb):
        sx = 19 + i * (sw + 3)
        rect(s, sx, cy, sw, sh, fill=PAPER, line=LINE_SOFT, line_w=0.5)
        rect(s, sx, cy, sw, 0.9, fill=GOLD)
        text_box(s, sx + 5, cy + 6, sw - 10, sh - 10, [
            ([t(f"STEP 0{i+1}", size=8, color=GOLD_T, serif=True, spc=1)],
             dict(after=3)),
            ([t(head, size=10.5, color=INK, serif=True, bold=True)],
             dict(after=2, spacing=1.2)),
            ([t(body, size=7.8, color=INK_SOFT)], dict(spacing=1.4)),
        ])
    ny = cy + sh + 8
    notes = [
        ("所要期間の目安",
         "お申し込みからプロフィール公開まで、2〜3週間程度を想定しています（審査状況により前後します）。"),
        ("先行登録のご案内",
         "正式リリースに先立ち、先行登録エージェント様を募集しています。初期に参画いただいた貴社エースのプロフィールは、リリース時点から求職者の比較・指名の対象となります。リリース時期・求職者側の集客計画は、ご説明の際に最新の状況をご案内します。"),
    ]
    for tag, body in notes:
        rect(s, 19, ny, 34, 6.6, line=LINE_GOLD, line_w=0.5)
        text_box(s, 19, ny + 1.6, 34, 4, [
            ([t(tag, size=7, color=GOLD_T, spc=1)], dict(align=PP_ALIGN.CENTER)),
        ])
        text_box(s, 58, ny + 0.6, 220, 12, [
            ([t(body, size=8, color=MUTED)], dict(spacing=1.4)),
        ])
        ny += 15
    footer(s, 12)

    # ---------- P.13 FAQ ----------
    s = add_slide(prs)
    y = header(s, "12", "FAQ", [[title_r("よくあるご質問。")]])
    faqs = [
        ("何名まで登録できますか？",
         "人数の上限ではなく「実績上位であること」が基準です。審査を通過したアドバイザーはご登録いただけます。まずは貴社のエース級の方からのお申し込みをおすすめします。"),
        ("費用はかかりますか？",
         "登録申請・審査は無料です。月額費用は通常5万円ですが、開始キャンペーンにより1年目は無料でご利用いただけます（2年目以降は月額5万円。詳細はP.11）。"),
        ("他社のアドバイザーと並ぶことに懸念があります。",
         "比較される場は、実力あるエースにとってはむしろ証明の場です。指名は求職者の意思で行われるため、実績と専門性が正当に評価されます。掲載内容の磨き込みも支援します。"),
        ("既存のスカウト媒体や自社集客と競合しませんか？",
         "MarkGateは「求職者からの指名」というプル型の接点で、既存のプッシュ型チャネルと役割が異なります。併用いただく前提の、追加チャネルとしてご活用ください。"),
        ("求職者の質はどのように担保されますか？",
         "求職者側も審査制です。経営幹部・管理職・高度専門職を中心とした、年収800万円以上を想定するハイクラス層に特化しています。"),
        ("求職者はどのように集めるのですか？",
         "審査制を軸に、ご紹介・招待、ハイクラス層向けの情報発信、提携チャネルなどを組み合わせ、質を担保しながら段階的に拡大する計画です。想定規模や進捗は、ご説明の際に最新の状況をご案内します。"),
    ]
    bw, bh, gap = 126.5, 36, 5
    for i, (q, a) in enumerate(faqs):
        bx = 19 + (i % 2) * (bw + gap)
        byy = y + 3 + (i // 2) * (bh + gap)
        rect(s, bx, byy, bw, bh, fill=PAPER, line=LINE_SOFT, line_w=0.5)
        text_box(s, bx + 6, byy + 4, bw - 12, bh - 8, [
            ([t("Q. ", size=9.5, color=GOLD_T, serif=True, bold=True),
              t(q, size=9.5, color=INK, serif=True, bold=True)], dict(after=2)),
            ([t("A. ", size=8, color=MUTED, serif=True),
              t(a, size=8, color=INK_SOFT)], dict(spacing=1.35)),
        ])
    footer(s, 13)

    # ---------- P.14 FAQ（情報の取り扱い・ご契約） ----------
    s = add_slide(prs)
    y = header(s, "13", "FAQ — DATA & CONTRACT",
               [[title_r("情報の取り扱い・ご契約について。")]])
    faqs2 = [
        ("MarkGateの法的な位置づけは？",
         "求職者と貴社アドバイザーの接点を提供するプラットフォームです。職業紹介（あっせん）および企業への紹介は、許可を保有する貴社の紹介事業として実施いただきます。当社も有料職業紹介事業許可の取得を申請準備中です。"),
        ("成約情報や求職者情報はどのように扱われますか？",
         "貴社の成約情報・求職者情報は、運営上必要な目的に限定して利用し、他のエージェント様への開示や競合比較への流用は行わない方針です。ご契約時に秘密保持条項を含む書面を締結します。"),
        ("掲載によって引き抜きリスクは高まりませんか？",
         "個人の直接連絡先は公開せず、指名はプラットフォームを経由する設計方針です。氏名・写真の掲載粒度も選択いただける想定です。確定済みの仕様・条件は、ご契約前にご説明します。"),
        ("アドバイザーが退職・異動した場合は？",
         "掲載は貴社とのご契約に紐づく設計です。退職・異動の際は掲載停止・所属情報の変更に対応し、プロフィールや指名履歴が個人アカウントとして他社へ引き継がれることはありません（詳細条件は契約書面にて確定します）。"),
    ]
    bw, bh, gap = 126.5, 48, 6
    for i, (q, a) in enumerate(faqs2):
        bx = 19 + (i % 2) * (bw + gap)
        byy = y + 4 + (i // 2) * (bh + gap)
        rect(s, bx, byy, bw, bh, fill=PAPER, line=LINE_SOFT, line_w=0.5)
        text_box(s, bx + 6, byy + 5, bw - 12, bh - 10, [
            ([t("Q. ", size=9.5, color=GOLD_T, serif=True, bold=True),
              t(q, size=9.5, color=INK, serif=True, bold=True)], dict(after=2.5)),
            ([t("A. ", size=8, color=MUTED, serif=True),
              t(a, size=8, color=INK_SOFT)], dict(spacing=1.35)),
        ])
    text_box(s, 19, y + 4 + 2 * (bh + gap) + 2, 259, 6, [
        ([t("※ 本ページの記載は現時点の設計方針です。確定した仕様・契約条件は、ご契約前に書面にてご案内します。",
            size=7.5, color=MUTED)], {}),
    ])
    footer(s, 14)

    # ---------- P.15 クロージング / 会社概要 ----------
    s = add_slide(prs, dark=True)
    text_box(s, 24, 42, 120, 6, [
        ([t("CONTACT", size=8.5, color=GOLD_L, spc=3)], {}),
    ])
    text_box(s, 24, 52, 130, 32, [
        ([t("まずは、貴社のエースを", size=25, color=WHITE, serif=True, bold=True)],
         dict(spacing=1.3)),
        ([t("一名から。", size=25, color=GOLD_L, serif=True, bold=True)],
         dict(spacing=1.3)),
    ])
    text_box(s, 24, 90, 128, 30, [
        ([t("トップアドバイザーだけが、開ける扉。", size=9.5,
            color=RGBColor.from_string("D5D9E2"))], dict(spacing=1.6)),
        ([t("その最初の一歩は、無料の登録申請と審査から。", size=9.5,
            color=RGBColor.from_string("D5D9E2"))], dict(spacing=1.6)),
        ([t("いまなら開始キャンペーンにより、1年目の月額費用も無料です。", size=9.5,
            color=RGBColor.from_string("D5D9E2"))], dict(spacing=1.6)),
        ([t("サービス詳細のご説明・デモのご案内も承ります。", size=9.5,
            color=RGBColor.from_string("D5D9E2"))], dict(spacing=1.6)),
    ])
    rect(s, 24, 126, 80, 12, fill=GOLD)
    text_box(s, 24, 129.2, 80, 6, [
        ([t("お問い合わせ・登録申請はこちら", size=9.5, color=NAVY, bold=True, spc=1)],
         dict(align=PP_ALIGN.CENTER)),
    ])
    text_box(s, 24, 141, 130, 16, [
        ([t("コーポレートサイトのお問い合わせフォームより「アドバイザー登録を申請したい」または「提携・協業について」をご選択ください。",
            size=7.5, color=DARK_SUB)], dict(spacing=1.5)),
        ([t("（サイトURLは確定後、本資料および配布用PDFに記載します）",
            size=7.5, color=DARK_SUB)], dict(spacing=1.5)),
    ])
    # 会社概要
    co = [
        ("会社名", "MarkGate株式会社（マークゲート）"),
        ("設立", "2026年7月13日"),
        ("代表取締役", "山本 朋鑑"),
        ("お問い合わせ", "コーポレートサイト お問い合わせフォームまで（URLは確定後に記載）"),
    ]
    cy = 58
    for k, v in co:
        text_box(s, 163, cy, 28, 10, [
            ([t(k, size=8.5, color=GOLD_L, spc=1)], {}),
        ])
        text_box(s, 195, cy, 82, 12, [
            ([t(v, size=8.5, color=RGBColor.from_string("DADDE4"))],
             dict(spacing=1.4)),
        ])
        cy += 13 if len(v) < 30 else 17
        rect(s, 163, cy - 3.5, 114, 0.3, fill=RGBColor.from_string("53492F"))
    text_box(s, 163, cy + 2, 114, 24, [
        ([t("※ 所在地・資本金などの詳細は、確定後に掲載いたします。",
            size=7, color=DARK_SUB)], dict(spacing=1.5)),
        ([t("※ 本資料の内容は2026年7月時点のものであり、予告なく変更される場合があります。",
            size=7, color=DARK_SUB)], dict(spacing=1.5)),
        ([t("※ 効果に関する記述はサービス設計に基づく想定であり、成果を保証するものではありません。",
            size=7, color=DARK_SUB)], dict(spacing=1.5)),
    ])
    footer(s, 15, dark=True)

    prs.save(out_path)
    print(f"saved: {out_path} ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "MarkGate_sales_deck.pptx"
    build(out)
